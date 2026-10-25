from __future__ import annotations

from collections.abc import Callable, Iterator
from pathlib import Path

import pytest
from PIL import Image

from lumen.cache import MemoryCache
from lumen.config import Settings
from lumen.embeddings import HashEmbedder
from lumen.engine import SearchEngine
from lumen.ingest.ocr import OcrPage

EngineFactory = Callable[..., SearchEngine]


SCANNED_SENTENCES = [
    "Passwords shall be changed after a suspected compromise.",
    "Terminals lock after ten minutes without input.",
    "Backup tapes are stored in a separate building.",
]


class FakeOcr:
    """Stands in for Tesseract, so OCR code paths run without the binary: page n reads as one sentence."""

    available = True
    languages = "eng+rus"

    def __init__(self) -> None:
        self.calls = 0

    def recognize(self, image: Image.Image) -> OcrPage:
        sentence = SCANNED_SENTENCES[self.calls % len(SCANNED_SENTENCES)]
        self.calls += 1
        return OcrPage(lines=["Scanned memo", "", sentence], confidence=91.5, words=9)


class NoOcr:
    available = False
    languages = "eng"

    def recognize(self, image: Image.Image) -> OcrPage:  # pragma: no cover - never called
        raise AssertionError("OCR is not available")


@pytest.fixture
def settings(tmp_path: Path) -> Settings:
    return Settings(
        data_dir=tmp_path / "data",
        embedding_model="hash",
        chunk_words=60,
        chunk_overlap_words=10,
        candidates=20,
        redis_url=None,
        api_key=None,
    )


@pytest.fixture
def make_engine(settings: Settings) -> Iterator[EngineFactory]:
    engines: list[SearchEngine] = []

    def make(**overrides: object) -> SearchEngine:
        options: dict[str, object] = {"embedder": HashEmbedder(256), "cache": MemoryCache(), "ocr": NoOcr()}
        options.update(overrides)
        engine = SearchEngine(settings, **options)  # type: ignore[arg-type]
        engines.append(engine)
        return engine

    yield make
    for engine in engines:
        engine.close()


@pytest.fixture
def engine(make_engine: EngineFactory) -> SearchEngine:
    return make_engine()


@pytest.fixture
def docs(tmp_path: Path) -> Path:
    """A folder with one small document of each kind (generated, so the tests show what they test)."""
    folder = tmp_path / "docs"
    folder.mkdir()
    build_docx(folder / "policy.docx")
    build_xlsx(folder / "controls.xlsx")
    build_pdf(folder / "guide.pdf")
    (folder / "notes.md").write_text(MARKDOWN, encoding="utf-8")
    (folder / "law.txt").write_text(RUSSIAN_TEXT, encoding="utf-8")
    return folder


MARKDOWN = """# Password Storage

Intro paragraph about [hashing](https://example.com) and **salts**.

## Argon2id

Use Argon2id with a minimum configuration of 19 MiB of memory, an iteration count of 2, and 1 degree of
parallelism.

```python
hash = argon2.hash(password)
```

| Algorithm | Memory |
| --- | --- |
| Argon2id | 19 MiB |

## bcrypt

- bcrypt has a maximum input length of 72 bytes.
- Use a work factor of 10 or more.
"""

RUSSIAN_TEXT = """Статья 5. Принципы обработки персональных данных

Хранение персональных данных должно осуществляться в форме, позволяющей определить субъекта
персональных данных, не дольше, чем этого требуют цели обработки персональных данных.

Статья 9. Согласие субъекта персональных данных

Субъект персональных данных принимает решение о предоставлении его персональных данных и дает согласие
на их обработку свободно, своей волей и в своем интересе.
"""


def build_docx(path: Path) -> None:
    import docx

    document = docx.Document()
    document.core_properties.title = "Access Policy"
    document.add_heading("1. Accounts", level=1)
    document.add_paragraph(
        "Accounts are disabled after 90 days without use. Owners review access every quarter."
    )
    document.add_heading("1.1 Administrators", level=2)
    document.add_paragraph("Administrator accounts require multi-factor authentication.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Role"
    table.cell(0, 1).text = "Review period"
    table.cell(1, 0).text = "Administrator"
    table.cell(1, 1).text = "Monthly"
    document.add_heading("2. Logging", level=1)
    document.add_paragraph("Audit logs are kept for one year and protected from modification.")
    document.save(str(path))


def build_xlsx(path: Path) -> None:
    import openpyxl

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Controls"
    sheet.append(["Control Identifier", "Control Name", "Control Text"])
    sheet.append(
        ["AC-7", "Unsuccessful Logon Attempts", "Enforce a limit of consecutive invalid logon attempts."]
    )
    sheet.append([None, None, None])  # an empty row keeps its Excel number for the rows after it
    sheet.append(
        ["SC-28", "Protection of Information at Rest", "Protect the confidentiality of information at rest."]
    )
    other = workbook.create_sheet("Notes")
    other.append(["Topic", "Note"])
    other.append(["Backups", "Backups are tested every month."])
    workbook.save(str(path))


def build_pdf(path: Path, *, pages: int = 3) -> None:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    pdf = canvas.Canvas(str(path), pagesize=A4)
    pdf.setTitle("Remote Access Guide")
    body = {
        1: [
            "1. Introduction",
            "This guide explains remote access for staff. It applies to laptops",
            "and phones.",
        ],
        2: [
            "2. Virtual Private Network",
            "Remote access requires the corporate VPN with multi-factor",
            "authentication.",
        ],
        3: [
            "3. Session Timeouts",
            "Idle sessions are closed after fifteen minutes of inactivity.",
            "Reconnect to continue.",
        ],
    }
    for number in range(1, pages + 1):
        pdf.setFont("Helvetica", 9)
        pdf.drawString(72, 800, "ACME Remote Access Guide - Internal")  # running header on every page
        pdf.drawString(290, 40, str(number))  # page number
        pdf.setFont("Helvetica", 11)
        y = 740
        for line in body.get(number, [f"Page {number} text about remote work."]):
            pdf.drawString(72, y, line)
            y -= 16
        pdf.showPage()
    pdf.save()


def build_scanned_pdf(path: Path) -> None:
    """An image-only PDF, like the output of a scanner."""
    image = Image.new("L", (1240, 1754), 255)
    image.save(str(path), "PDF", resolution=150)


@pytest.fixture
def scanned_pdf(tmp_path: Path) -> Path:
    path = tmp_path / "scan.pdf"
    build_scanned_pdf(path)
    return path
