"""Download the demo corpus from its official sources and rebuild the derived files.

    python corpus/scripts/fetch.py            # download into corpus/files/ and check the checksums
    python corpus/scripts/fetch.py --verify   # only check the committed files

Every document is either a U.S. government work (not subject to copyright in the United States), an
official Russian state document (not an object of copyright under art. 1259(6) of the Civil Code of the
Russian Federation), or an OWASP Cheat Sheet (CC BY-SA 4.0). See corpus/README.md.

Some Russian government sites refuse connections from outside Russia, which is why the files are
committed and CI only verifies them.
"""

from __future__ import annotations

import argparse
import datetime as dt
import hashlib
import html
import io
import json
import sys
import urllib.request
import zipfile
from html.parser import HTMLParser
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
FILES = ROOT / "files"
MANIFEST = ROOT / "manifest.json"

# The consolidated text of 152-FZ from the official legal information system (pravo.gov.ru).
LAW_152_HTML = "http://pravo.gov.ru/proxy/ips/?doc_itself=&nd=102108261&page=1&rdk=38"


class _LawParagraphs(HTMLParser):
    """Paragraphs of a pravo.gov.ru document: ``<p class="H">`` marks chapter and article headings.

    Superscripts (``<span class="W9">``) number inserted articles and parts: "Статья 10<sup>1</sup>" is
    cited as "статья 10.1", so the superscript becomes ".1".
    """

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.paragraphs: list[tuple[str, str]] = []
        self._class: str | None = None
        self._text: list[str] = []

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag == "p":
            self._class = dict(attrs).get("class") or ""
            self._text = []
        elif tag == "span" and dict(attrs).get("class") == "W9" and self._class is not None:
            self._text.append(".")

    def handle_endtag(self, tag: str) -> None:
        if tag == "p" and self._class is not None:
            text = " ".join(html.unescape("".join(self._text)).replace("\xa0", " ").split())
            if text:
                self.paragraphs.append((self._class, text))
            self._class = None

    def handle_data(self, data: str) -> None:
        if self._class is not None:
            self._text.append(data)


def build_law_docx(source_html: bytes, target: Path) -> None:
    """Rebuild the Word copy of 152-FZ: chapters as Heading 1, articles as Heading 2, text unchanged."""
    import docx

    parser = _LawParagraphs()
    parser.feed(source_html.decode("cp1251"))
    document = docx.Document()
    properties = document.core_properties
    properties.title = "Федеральный закон от 27.07.2006 № 152-ФЗ «О персональных данных»"
    properties.subject = "Действующая редакция, текст с официального портала pravo.gov.ru"
    properties.author = ""
    properties.comments = LAW_152_HTML
    properties.language = "ru-RU"
    for css_class, text in parser.paragraphs:
        if css_class == "H" and text.startswith("Глава"):
            document.add_heading(text, level=1)
        elif css_class == "H" and text.startswith("Статья"):
            document.add_heading(text, level=2)
        else:
            document.add_paragraph(text)
    # python-docx stamps the current time, in the properties and in the ZIP entries; fix both so a
    # rebuild is byte-for-byte identical and the checksum in the manifest stays valid
    stamp = dt.datetime(2026, 1, 1, tzinfo=dt.timezone.utc)
    properties.created = properties.modified = properties.last_printed = stamp
    properties.revision = 1
    buffer = io.BytesIO()
    document.save(buffer)
    with zipfile.ZipFile(buffer) as source, zipfile.ZipFile(target, "w", zipfile.ZIP_DEFLATED) as out:
        for item in source.infolist():
            fixed = zipfile.ZipInfo(item.filename, date_time=(2026, 1, 1, 0, 0, 0))
            fixed.compress_type = zipfile.ZIP_DEFLATED
            out.writestr(fixed, source.read(item.filename))


def download(url: str) -> bytes:
    request = urllib.request.Request(url, headers={"User-Agent": "lumen-search corpus fetch"})
    with urllib.request.urlopen(request, timeout=120) as response:  # fixed, known URLs
        return bytes(response.read())


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--verify", action="store_true", help="only check the committed files")
    args = parser.parse_args()
    manifest = json.loads(MANIFEST.read_text(encoding="utf-8"))
    FILES.mkdir(exist_ok=True)
    failed = 0
    for entry in manifest["documents"]:
        target = ROOT / entry["file"]
        if not args.verify:
            if entry["id"] == "ru-152-fz-personal-data":
                build_law_docx(download(LAW_152_HTML), target)
            else:
                target.write_bytes(download(entry["download_url"]))
        if not target.exists():
            print(f"missing  {entry['file']}")
            failed += 1
            continue
        digest = sha256(target)
        expected = entry.get("sha256")
        status = "ok" if digest == expected else ("NEW" if expected is None else "CHANGED")
        if status != "ok" and args.verify:
            failed += 1
        print(f"{status:8} {entry['file']}  {digest[:16]}")
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())
