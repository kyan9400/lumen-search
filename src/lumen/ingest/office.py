"""Word (.docx), Excel (.xlsx) and CSV parsing."""

from __future__ import annotations

import csv
import datetime as dt
import re
from collections.abc import Iterable, Iterator, Sequence
from pathlib import Path
from typing import Any

from lumen.ingest.errors import IngestError
from lumen.ingest.layout import detect_heading
from lumen.models import Block, ParsedDocument
from lumen.textproc import detect_language

_HEADING_STYLE = re.compile(r"(?:heading|заголовок)\s*(\d)", re.IGNORECASE)
_TITLE_STYLES = {"title", "название", "заголовок"}
MAX_ROWS = 100_000


def parse_docx(path: Path) -> ParsedDocument:
    import docx
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    try:
        document = docx.Document(str(path))
    except Exception as exc:  # python-docx raises several unrelated types for broken packages
        raise IngestError(f"cannot open Word document: {exc}") from exc

    def heading_level(paragraph: Paragraph) -> int | None:
        style = paragraph.style
        while style is not None:
            name = (style.name or "").strip().lower()
            match = _HEADING_STYLE.match(name)
            if match:
                return int(match.group(1))
            if name in _TITLE_STYLES:
                return 1
            style = style.base_style
        ppr = paragraph._p.pPr
        outline = ppr.find(qn("w:outlineLvl")) if ppr is not None else None
        if outline is not None:
            level = int(outline.get(qn("w:val"), "9"))
            return level + 1 if level < 9 else None
        return None

    blocks: list[Block] = []
    first_heading = None
    for item in document.iter_inner_content():
        if isinstance(item, Paragraph):
            text = item.text.strip()
            if not text:
                continue
            level = heading_level(item)
            if level is None and (detected := detect_heading(text)):
                level = detected.level  # many documents format headings by hand instead of with styles
            if level is not None:
                first_heading = first_heading or text
                blocks.append(Block(text, kind="heading", level=level))
            else:
                blocks.append(Block(text))
        elif isinstance(item, Table):
            for row in item.rows:
                cells: list[str] = []
                for cell in row.cells:  # a merged cell is returned once per grid column it spans
                    value = " ".join(cell.text.split())
                    if value and (not cells or cells[-1] != value):
                        cells.append(value)
                if cells:
                    blocks.append(Block(" | ".join(cells), kind="table"))
    if not blocks:
        raise IngestError("the Word document has no text")
    title = (document.core_properties.title or "").strip() or first_heading or path.stem
    return ParsedDocument(
        title=title,
        media_type="docx",
        blocks=blocks,
        language=detect_language(" ".join(b.text for b in blocks[:400])),
    )


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, dt.datetime):
        return value.date().isoformat() if value.time() == dt.time() else value.isoformat(sep=" ")
    if isinstance(value, dt.date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return " ".join(str(value).split())


def _header_index(rows: Sequence[Sequence[str]]) -> int | None:
    """The first of the opening rows that reads like column headers (two or more text cells)."""
    for index, row in enumerate(rows[:10]):
        filled = [cell for cell in row if cell]
        textual = [cell for cell in filled if not re.fullmatch(r"[\d.,:%\- ]+", cell)]
        if len(filled) >= 2 and len(textual) >= 0.7 * len(filled):
            return index
    return None


def _records(sheet: str, rows: Iterable[Sequence[str]], first_row_number: int = 1) -> Iterator[Block]:
    """One block per data row: ``Header: value`` lines, labelled by the first two columns."""
    buffered: list[tuple[int, list[str]]] = []
    iterator = iter(rows)
    for number, row in enumerate(iterator, start=first_row_number):
        buffered.append((number, list(row)))
        if len(buffered) == 10:
            break
    header_at = _header_index([row for _, row in buffered])
    if header_at is None:
        width = max((len(row) for _, row in buffered), default=0)
        header = [f"Column {i + 1}" for i in range(width)]
        data: Iterable[tuple[int, list[str]]] = buffered
    else:
        header = [cell or f"Column {i + 1}" for i, cell in enumerate(buffered[header_at][1])]
        data = buffered[header_at + 1 :]

    def all_rows() -> Iterator[tuple[int, list[str]]]:
        yield from data
        for number, row in enumerate(iterator, start=first_row_number + len(buffered)):
            yield number, list(row)

    for count, (number, row) in enumerate(all_rows()):
        if count >= MAX_ROWS:
            break
        pairs = [
            (header[i] if i < len(header) else f"Column {i + 1}", value)
            for i, value in enumerate(row)
            if value
        ]
        if not pairs:
            continue
        label_parts = [value for value in row[:2] if value and len(value) <= 120]
        yield Block(
            "\n".join(f"{name}: {value}" for name, value in pairs),
            kind="row",
            sheet=sheet,
            row=number,
            label=" ".join(label_parts) or None,
        )


def parse_xlsx(path: Path) -> ParsedDocument:
    import openpyxl

    try:
        workbook = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    except Exception as exc:  # zipfile, KeyError and InvalidFileException all signal a bad workbook
        raise IngestError(f"cannot open Excel workbook: {exc}") from exc
    try:
        blocks: list[Block] = []
        for sheet in workbook.worksheets:
            rows = ([_cell_text(v) for v in row] for row in sheet.iter_rows(values_only=True))
            blocks.extend(_records(sheet.title, rows))
        title = (workbook.properties.title or "").strip() or path.stem
    finally:
        workbook.close()
    if not blocks:
        raise IngestError("the workbook has no data rows")
    return ParsedDocument(
        title=title,
        media_type="xlsx",
        blocks=blocks,
        language=detect_language(" ".join(b.text for b in blocks[:400])),
    )


def parse_csv(path: Path) -> ParsedDocument:
    raw = path.read_bytes()
    try:
        text = raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = raw.decode("cp1251")  # Excel on Russian Windows saves CSV in Windows-1251
    try:
        dialect: type[csv.Dialect] | csv.Dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    rows = ([" ".join(cell.split()) for cell in row] for row in csv.reader(text.splitlines(), dialect))
    blocks = list(_records(path.stem, rows))
    if not blocks:
        raise IngestError("the CSV file has no data rows")
    return ParsedDocument(
        title=path.stem,
        media_type="csv",
        blocks=blocks,
        language=detect_language(" ".join(b.text for b in blocks[:400])),
    )
